"""Hana Residence quotation audit workbook (R3, after client answers 11 Oct). Run: python3 build_audit_r3.py <out.xlsx>"""
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT = sys.argv[1]
CURRENT_TOTAL = 1043267.62
PREV_TOTAL = 971295.16  # sheet total before R3 edits (R2 audit showed 990,075.16; curtains 16,280 + protection 2,500 were removed in the sheet afterwards)

HDR = PatternFill("solid", fgColor="3B3A36")
HFONT = Font(bold=True, color="FFFFFF")
RED = PatternFill("solid", fgColor="F8D7D3")
AMB = PatternFill("solid", fgColor="FCEBC7")
GRN = PatternFill("solid", fgColor="DDEFD9")
THIN = Side(style="thin", color="BBBBBB")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
PRIO_FILL = {"A": RED, "B": AMB, "C": GRN}

wb = Workbook()


def sheet(title, headers, widths):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
        c = ws.cell(row=1, column=i)
        c.fill, c.font = HDR, HFONT
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "A2"
    return ws


def finish(ws, money_cols=(), prio_col=None):
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.border = BOX
            if c.column in money_cols and isinstance(c.value, (int, float)):
                c.number_format = '#,##0.00'
        if prio_col:
            p = row[prio_col - 1].value
            if p in PRIO_FILL:
                row[prio_col - 1].fill = PRIO_FILL[p]


# ---------------------------------------------------------------- 1. Missing / add items (still open after R3)
# (section, item, qty, unit, rate, basis / where seen, priority)
ADD = [
    ("Curtains", "Curtains & blinds – section now RM0 (sheer 34 ft has no rate; all other rooms deleted after R2). R1/R2 had RM16,280: sheer 34ft@160, kitchen blinds, double-layer Parents/Master/Son, roman WIC/Daughter, laundry/maid blinds",
     1, "L/S", 16280.00, "If curtains are now by another vendor, add 'Curtains by owner' to Exclusions instead. Do not leave a priced line at RM0.", "A"),
    ("Preliminaries", "Floor / lift / staircase protection (removed from sheet after R2, was RM2,500)", 1, "L/S", 2500.00,
     "Protection of retained tiles in Son's & Daughter's bathrooms, existing staircase & glass balustrade (kept) is now more important, not less.", "A"),
    ("Wet Works", "PC Sum – tile MATERIAL: Powder, Master Bath, FF/Parents Bath, Entrance, Laundry floor+wall, Wet Kitchen wall, Terraces (approx. 2,500 sqft incl. 10% wastage @ PC RM12/sqft)", 2500, "sqft", 12.00,
     "Quote still says 'TBC upon final selection' = RM0. Reduced from 3,600 sqft because Son's & Daughter's bathrooms keep existing tiles; terrace 315 sqft added.", "A"),
    ("Wet Works", "Laundry wall tiles – Dong Peng Limestone JH12352GZ 600x1200 structured (labour)", 180, "sqft", 35.00,
     "GF legend red wall-tile lines along laundry walls. 3D laundry shows full-height wall tile.", "A"),
    ("Wet Works", "Wet kitchen wall tile – quantity top-up (take-off ~330 vs 250 sqft)", 80, "sqft", 35.00, "", "B"),
    ("Wet Works", "Staircase risers & nosing (Porter covers treads only) – PC sum", 1, "L/S", 3500.00, "Still open with Porter.", "A"),
    ("Wet Works", "Removal of existing timber flooring 1,647 sqft – now written into Demolition L/S description at no extra cost", 1647, "sqft", 2.86,
     "Confirm the demolition subcon's RM37,145 covers it. If not, cost ~RM2/sqft -> sell RM2.86 (+RM4.7k).", "B"),
    ("Ceiling", "Ceiling quantity top-up (corridor & stair, powder, utility, maid + under-measured master, WIC, son, daughter)", 1000, "sqft", 6.00,
     "Quote bedroom ceiling 2,799 sqft vs rough take-off ~3,800.", "B"),
    ("Ceiling", "Moisture-resistant ceiling – Parents Bath & Powder", 70, "sqft", 7.00, "Son's & Daughter's done in R3.", "A"),
    ("Ceiling", "AC grille cut-out / access panels (only if ducted / cassette units chosen)", 10, "nos", 150.00, "Aircond now quoted as wall-mounted PS – confirm type.", "C"),
    ("Ceiling", "Curved arch / curved wall-to-ceiling transition at staircase & corridor", 2, "nos", 1200.00, "3D p.4-7, p.37.", "B"),
    ("Preliminaries", "Scaffolding for double-volume void", 1, "L/S", 3500.00, "Living/dining void ~6-7 m.", "A"),
    ("Preliminaries", "Final chemical cleaning, approx. 5,500 sqft", 1, "L/S", 3800.00, "", "A"),
    ("Preliminaries", "Contractor All-Risk insurance", 1, "L/S", 2500.00, "Optional.", "C"),
    ("Joinery", "Foyer / Living divider c/w planter box & concealed lighting", 1, "nos", 8500.00, "3D p.8-12, plan label 'DIVIDER'.", "A"),
    ("Joinery", "Second hidden door @ Utility / Powder", 1, "nos", 4500.00, "Plan labels 'HIDDEN DOOR' x2.", "A"),
    ("Joinery", "Maid room tatami platform c/w storage", 1, "nos", 2800.00, "", "B"),
    ("Joinery", "Parents Bath – floating vanity + tall storage cabinet", 1, "L/S", 6000.00, "3D p.43. Depends on 'First Floor Bathroom' = Parents Bath.", "A"),
    ("Joinery", "Parents Bath – mirror cabinet c/w LED", 1, "nos", 1200.00, "Son's & Daughter's done in R3.", "A"),
    ("Joinery", "Powder room – vanity / stone basin base & mirror", 1, "L/S", 3000.00, "3D p.44.", "B"),
    ("Joinery", "Staircase side wall – full-height V-groove panelling", 1, "L/S", 8500.00, "3D p.37-38.", "B"),
    ("Joinery", "Master Bath – open-shelf tall storage", 1, "nos", 3500.00, "", "C"),
    ("Table Top", "Dry Kitchen island top-up for 900 deep + waterfall ends", 1, "L/S", 3400.00, "", "A"),
    ("Table Top", "Parents Bath vanity top + open hole", 1, "L/S", 1270.00, "4 ft @ 280 + 1 hole @ 150.", "A"),
    ("Doors & Windows", "10mm tempered glass shower screen – Parents Bath", 1, "nos", 1800.00, "Son's & Daughter's keep existing.", "A"),
    ("Doors & Windows", "Existing doors – refurbish / laminate c/w new black lever handles (~10 leaves)", 10, "nos", 650.00, "", "B"),
    ("Electrical", "32A dedicated point for induction hob (hob by owner, point by Groo)", 1, "nos", 650.00, "", "A"),
    ("Electrical", "Additional 15A points – water heaters, dishwasher, wine chiller", 4, "nos", 350.00, "", "A"),
    ("Electrical", "PROVISIONAL – DB upgrade / sub-DB (smart home + aircond + heavy appliances)", 1, "L/S", 4500.00, "Aircond now in scope -> load increases.", "A"),
    ("Electrical", "Aircond power points – top-up to match 6 units in PS (4 quoted)", 2, "nos", 300.00, "", "B"),
    ("Electrical", "LED step-light points @ staircase", 10, "nos", 160.00, "", "B"),
    ("Electrical", "Exhaust fan points – powder & bathrooms", 4, "nos", 200.00, "", "C"),
    ("Plumbing", "PROVISIONAL – piping re-route for new vanity positions (Parents, Son's, Daughter's)", 3, "nos", 1200.00,
     "New vanity cabinets may not match existing inlet/outlet positions.", "B"),
    ("Curtains", "Family Area – full-height sheer + night curtain", 14, "ft", 160.00, "3D p.41-44.", "B"),
]

ws = sheet("1 Missing Items", ["#", "Section", "Item to add", "Qty", "Unit", "Suggested sell rate (RM)", "Amount (RM)", "Why / where seen", "Priority"],
           [4, 14, 60, 8, 7, 12, 13, 60, 8])
for i, (sec, item, q, u, r, why, p) in enumerate(ADD, 1):
    ws.append([i, sec, item, q, u, r, round(q * r, 2), why, p])
n = len(ADD) + 1
ws.append([])
for lbl, f in (("TOTAL – Priority A", f'=SUMIF(I2:I{n},"A",G2:G{n})'), ("TOTAL – Priority B", f'=SUMIF(I2:I{n},"B",G2:G{n})'),
               ("TOTAL – Priority C", f'=SUMIF(I2:I{n},"C",G2:G{n})'), ("TOTAL – all", f"=SUM(G2:G{n})")):
    ws.append(["", "", lbl, "", "", "", f])
for r in range(n + 2, n + 6):
    ws.cell(row=r, column=3).font = Font(bold=True)
    ws.cell(row=r, column=7).font = Font(bold=True)
    ws.cell(row=r, column=7).number_format = '#,##0.00'
finish(ws, money_cols=(6, 7), prio_col=9)

add_A = sum(q * r for _, _, q, _, r, _, p in ADD if p == "A")
add_B = sum(q * r for _, _, q, _, r, _, p in ADD if p == "B")
add_C = sum(q * r for _, _, q, _, r, _, p in ADD if p == "C")

# ---------------------------------------------------------------- 2. R3 changes made in the Google Sheet
CHG = [
    ("Header", "F6 date", "10/8/2026", "08 Oct 2026", 0),
    ("Header", "J8 company address", "143 Jln Desa Utama, Taman Danau Desa", "Level 3A, Wisma JAG, Jalan Desa Utama, Taman Desa, 58100 KL", 0),
    ("1.0 Ceiling", "Son's & Daughter's Bathroom – dismantle + 9mm MR ceiling c/w access panel (new rows 84-85)", "-", "162 sqft @ RM9", 1458),
    ("2.0 Painting", "Ceiling paint qty now includes bathroom ceiling", "2,799", "2,961 sqft", 162),
    ("3.0 Wet Works", "Son's Bathroom tiling, floor trap, mitre – deleted (tiles remain)", "15,324.14", "0", -15324.14),
    ("3.0 Wet Works", "Daughter's Bathroom tiling, floor trap, mitre – deleted (tiles remain)", "14,407.72", "0", -14407.72),
    ("3.0 Wet Works", "Terrace 1, 2 & 3 – outdoor anti-slip tiles c/w screed to fall, waterproofing, ponding test (labour; tile material TBC)", "-", "315 sqft @ RM45 (cost 31.5/0.7)", 14175),
    ("3.0 Wet Works", "Screed -> cement screed + self-levelling for glue-down (Porter 20-yr warranty)", "1,647 @ RM5", "1,647 @ RM8", 4941),
    ("3.0 Wet Works", "Demolition L/S description – adds removal of existing timber flooring; bathroom scope changed (amount unchanged)", "53,064.29", "53,064.29", 0),
    ("4.0 Plumbing", "Water heater installation @ Master & First Floor Bath (heater by owner)", "-", "2 @ RM250", 500),
    ("4.0 Plumbing", "Descriptions: sinks/taps/appliances & water heaters by owner; bathroom reinstall line covers basin re-fix to new vanity", "-", "-", 0),
    ("7.0 Doors & Windows", "Wet Kitchen Option 2 – hack & enlarge sink-wall window c/w new black alu window", "-", "L/S RM4,500", 4500),
    ("8.0 Joinery", "Son's & Daughter's Bath – floating vanity sink cabinet x2 + mirror cabinet c/w LED x2", "-", "2 @ 2,600 + 2 @ 1,200", 7600),
    ("8.0 Joinery", "Master Bedroom – TV wall reworded to Option 2 (twin round columns + low fireplace plinth)", "13,500", "13,500", 0),
    ("9.0 Table Top", "Son's & Daughter's vanity top 6 ft @ 280 + 2 more open holes", "-", "1,680 + 300", 1980),
    ("12.0 Smart Home", "15% markup on all 5 lines; supplier name removed from client note", "142,588.80", "163,977.12", 21388.32),
    ("13.0 Air-Cond", "PROVISIONAL SUM – 6 inverter wall-mounted units c/w piping @ RM5,000", "Not included", "30,000", 30000),
    ("14.0 Landscape", "PROVISIONAL SUM – front garden & terraces", "-", "15,000", 15000),
    ("Exclusions", "New block under Grand Total: loose furniture, appliances, sanitary, taps & sinks, water heaters, decorative lights, fireplace, artwork; gate; JMB deposit/permits", "-", "-", 0),
    ("Payment terms", "Smart Home billed separately: 50% deposit / 40% delivery / 10% T&C", "-", "-", 0),
]
ws = sheet("2 R3 Changes", ["Section", "Change", "Before", "After", "Δ RM"], [16, 80, 22, 30, 13])
for row in CHG:
    ws.append(list(row))
ws.append(["", "NET CHANGE", "", "", f"=SUM(E2:E{len(CHG)+1})"])
ws.cell(row=len(CHG) + 2, column=2).font = Font(bold=True)
finish(ws, money_cols=(5,))

# ---------------------------------------------------------------- 3. Margin estimate by section (R3)
MARGIN = [
    ("Preliminaries", 2100, 1500), ("1.0 Ceiling", 38659, 27100), ("2.0 Painting", 24761, 17100),
    ("3.0 Demolition & Masonry / Wet Works", 247447.5, 176850), ("4.0 Plumbing", 9850, 6850),
    ("5.0 Electrical", 37470, 25000), ("6.0 Light Fittings & Fans", 2100, 1500), ("7.0 Doors & Windows", 17400, 12700),
    ("8.0 Joinery", 339250, 215000), ("9.0 Table Top", 23120, 16400), ("10.0 Flooring (Porter, incl. staircase)", 82233, 76259),
    ("11.0 Curtains & Blinds", 0, 0), ("12.0 Smart Home (15%)", 163977.12, 142588.80), ("13.0 Air-Cond (PS)", 30000, 25500),
    ("14.0 Landscape (PS)", 15000, 12000), ("Project Management", 9900, 5000),
]
ws = sheet("3 Margin Estimate", ["Section", "Quoted (RM)", "Est. cost (RM)", "Gross profit (RM)", "GP %"], [36, 15, 15, 16, 9])
for i, (s, sell, cost) in enumerate(MARGIN, 2):
    ws.append([s, sell, cost, f"=B{i}-C{i}", f"=IF(B{i}=0,0,D{i}/B{i})"])
last = len(MARGIN) + 1
sh = 2 + [m[0] for m in MARGIN].index("12.0 Smart Home (15%)")
ws.append(["TOTAL", f"=SUM(B2:B{last})", f"=SUM(C2:C{last})", f"=B{last+1}-C{last+1}", f"=D{last+1}/B{last+1}"])
ws.append(["TOTAL excl. Smart Home", f"=B{last+1}-B{sh}", f"=C{last+1}-C{sh}", f"=B{last+2}-C{last+2}", f"=D{last+2}/B{last+2}"])
for r in range(2, last + 3):
    for c in (2, 3, 4):
        ws.cell(row=r, column=c).number_format = '#,##0'
    ws.cell(row=r, column=5).number_format = '0.0%'
for r in (last + 1, last + 2):
    for c in range(1, 6):
        ws.cell(row=r, column=c).font = Font(bold=True)
ws.append([])
ws.append(["Est. cost from QS Pricing Matrix ranges + supplier quotes (Dong Peng, Porter, M Vida). Excludes Porter's 10% designer commission (~RM6.4k, off-quote). Aircond / landscape costs assume ~15-20% on PS."])
finish(ws)
m_sell = sum(m[1] for m in MARGIN); m_cost = sum(m[2] for m in MARGIN)
sm = [m for m in MARGIN if m[0].startswith("12.0")][0]
gp = (m_sell - m_cost) / m_sell
gp_ex = ((m_sell - sm[1]) - (m_cost - sm[2])) / (m_sell - sm[1])

# ---------------------------------------------------------------- 4. Rate / margin flags still open
RATE = [
    ("Project Management", "PM fee RM1.80/sqft = RM9,900", "~0.9% of RM1.04M", "Min 3% (~RM31k)", "Now also coordinating aircond, landscape & smart home vendors.", "A"),
    ("Flooring", "Porter S&I RM39/sqft", "~6% + 10% commission", "25%", "Thin but low-risk. Negotiate install RM4.5.", "B"),
    ("Flooring", "Staircase steps RM18,000", "~11% + commission", "25%", "Risers & nosing still excluded.", "B"),
    ("Ceiling", "9mm plaster RM6/sqft, L-box RM28/ft", "17-40%", "45%", "Matrix sell RM10 / RM50.", "B"),
    ("Wet Works", "Decimal rates (37.14, 53,064.29, 45.00 ...)", "30%", "-", "Decimals reveal cost ÷ 0.7. Round before sending.", "B"),
    ("Air-Cond", "PS 6 units @ RM5,000", "~15% assumed", "15-20%", "Replace with real supplier quote + layout (units, HP, wall vs cassette/ducted).", "A"),
    ("Landscape", "PS RM15,000", "~20% assumed", "20-25%", "Replace with landscape subcon quote once scope agreed.", "A"),
    ("Wet Works", "Self-levelling RM8/sqft", "~35%", "30%", "Porter quoted self-levelling RM3-4 + screed. OK.", "C"),
]
ws = sheet("4 Rate Flags", ["Section", "Item", "Est. margin", "Target", "Comment", "Priority"], [16, 40, 18, 16, 70, 8])
for row in RATE:
    ws.append(list(row))
finish(ws, prio_col=6)

# ---------------------------------------------------------------- 5. Client answers (11 Oct) + still-open questions
ANS = [
    ("Answered", "Son's & Daughter's bathrooms", "Keep tiles; only mirror cabinet + sink cabinet + ceiling redo", "Applied (tiling deleted, ceiling/joinery/table top added)"),
    ("Answered", "Maid Bath", "No works", "Nothing in quote – OK"),
    ("Answered", "Terraces after decking removal", "Tiles floor", "Applied – 315 sqft labour; material in tile PC (Missing #3)"),
    ("Answered", "Existing glass balustrade", "Keep", "Balustrade removed from missing list; protection added to Missing #2"),
    ("Answered", "Wet Kitchen / Master Bedroom option", "Wet Kitchen Opt 2, Master Bedroom Opt 2", "Applied – see OPEN #3 on Wet Kitchen Opt 2 reading"),
    ("Answered", "Sinks/taps, hood/hob/oven, water heaters, fireplace", "Owner", "Applied – exclusions + install labour only"),
    ("Answered", "Smart home", "Groo quotes it in", "Applied – 15% markup (+RM21,388.32)"),
    ("Answered", "Unit size", "5,500 sqft", "Kept"),
    ("Answered", "Company address", "Changed", "Applied – J8"),
    ("Answered", "Quote date", "8 Oct", "Applied – '08 Oct 2026'"),
    ("Answered", "Design fee", "Billed separately, done", "No change"),
    ("Answered", "Smart-home deposit", "Most likely follow", "Applied – 50/40/10 line under payment terms"),
    ("Answered", "Exclusions", "Aircond & landscape quoted in; others by owner; no gate; JMB deposit by owner", "Applied – sections 13 & 14 (PS) + Exclusions block"),
    ("OPEN", "1. 'First Floor Bathroom' (RM13,419) – FF only has Master, Son's, Daughter's. Is it the Parents Bath (GF)?", "", "If yes rename + add Parents Bath joinery/top/screen (Missing #18-19, #24-25)"),
    ("OPEN", "2. Existing wood floor -> dismantle + self-levelling (applied, 20-yr warranty). Reply to Porter & confirm demolition subcon covers removal.", "", "Revert: row 114 RM8 -> RM5 (-RM4,941)"),
    ("OPEN", "3. Wet Kitchen Opt 2 read from 3D p.33 as: larger black-frame window at sink wall, rest same as Opt 1 (curved island, reeded sliding doors kept).", "", "If Opt 2 also drops island / sliding doors -> deduct rows 180, 212, Table Top 268"),
    ("OPEN", "4. Curtains section RM0 + protection removed after R2 – intentional?", "", "Missing #1-2 (RM18,780)"),
    ("OPEN", "5. Aircond: no. of units, wall-mounted vs cassette/ducted, brand.", "", "Replace PS"),
    ("OPEN", "6. Landscape scope / designer.", "", "Replace PS"),
    ("OPEN", "7. Foyer finish: Ling Hua Bai (Dong Peng count) vs Bellezza (DWG legend).", "", ""),
    ("OPEN", "8. Porter: area breakdown of 1,647 sqft (drawing ~1,880) + staircase risers/nosing price.", "", ""),
]
ws = sheet("5 Questions", ["Status", "Question", "Answer", "Action in sheet"], [10, 80, 45, 55])
for row in ANS:
    ws.append(list(row))
finish(ws)
for row in ws.iter_rows(min_row=2):
    row[0].fill = GRN if row[0].value == "Answered" else AMB

# ---------------------------------------------------------------- 6. Supplier quotes
SUP = [
    ("Dong Peng (DP Ceramic)", "09/10/2026", "Ling Hua Bai 900x1800x9 honed, 91 pcs @ RM250", 22750, "30,030 (91 @ RM330)", "Logistics TBC, unloading excl., 50% deposit, 12-wk production"),
    ("Porter Rosewood", "09/10/2026", "Art Deco E30 straight plank 1,647 sqft", 51200, "64,233 incl. install", "Incl. 10% designer commission; area breakdown TBC"),
    ("Porter Rosewood (installer est.)", "10/10/2026", "Glue-down install 1,647 sqft @ RM5.5", 9059, "(in RM39/sqft)", "20-yr warranty only on self-levelled cement / tile"),
    ("Porter Rosewood", "09/10/2026", "Staircase steps 1.2mm + install RM3k", 16000, "18,000", "Steps only"),
    ("M Vida (AETHO)", "13/09/2026", "Smart home package", 142588.80, "163,977.12 (15%)", "3-month lead time; 50/40/10 to client"),
    ("Aircond – TBC", "-", "6 units PS", None, "30,000 PS", "Get quote"),
    ("Landscape – TBC", "-", "PS", None, "15,000 PS", "Get quote"),
]
ws = sheet("6 Supplier Quotes", ["Supplier", "Date", "Item", "Cost (RM)", "Quoted to client", "Notes"], [26, 11, 46, 13, 24, 55])
for row in SUP:
    ws.append(list(row))
finish(ws, money_cols=(4,))

# ---------------------------------------------------------------- 0. Summary
ws = wb.active
ws.title = "0 Summary"
ws.column_dimensions["A"].width = 72
ws.column_dimensions["B"].width = 20
rows = [
    ("HANA RESIDENCE – QUOTATION AUDIT R3 (client answers applied, 11 Oct 2026)", None),
    ("Quotation 261015 Grand Total after R3", CURRENT_TOTAL),
    ("Sheet total before R3 edits", PREV_TOTAL),
    ("Net change R3 (see sheet 2)", round(CURRENT_TOTAL - PREV_TOTAL, 2)),
    ("Est. gross margin R3 (sheet 3, excl. Porter commission)", f"≈{gp*100:.1f}%  (≈{gp_ex*100:.1f}% excl. smart home)"),
    ("", None),
    ("A. Still missing – Priority A", round(add_A, 2)),
    ("B. Still missing – Priority B", round(add_B, 2)),
    ("C. Optional – Priority C", round(add_C, 2)),
    ("", None),
    ("Revised total – R3 + A", round(CURRENT_TOTAL + add_A, 2)),
    ("Revised total – R3 + A + B", round(CURRENT_TOTAL + add_A + add_B, 2)),
    ("", None),
    ("Backups: 'Hana Sheon & Tee - BACKUP R1 (09 Oct)' and 'Hana Sheon & Tee - BACKUP R2 (11 Oct)' (state just before R3 edits).", None),
]
for a, b in rows:
    ws.append([a, b])
ws["A1"].font = Font(bold=True, size=13)
for r in range(2, 14):
    if isinstance(ws.cell(row=r, column=2).value, (int, float)):
        ws.cell(row=r, column=2).number_format = '#,##0.00'
for r in (11, 12):
    ws.cell(row=r, column=1).font = Font(bold=True)
    ws.cell(row=r, column=2).font = Font(bold=True)
ws["A14"].alignment = Alignment(wrap_text=True)

wb.save(OUT)
print(f"A={add_A:,.2f} B={add_B:,.2f} C={add_C:,.2f}  GP={gp:.3f} ex={gp_ex:.3f}  sell={m_sell:,.2f}")
print(f"R3+A={CURRENT_TOTAL+add_A:,.2f}  R3+A+B={CURRENT_TOTAL+add_A+add_B:,.2f}")
